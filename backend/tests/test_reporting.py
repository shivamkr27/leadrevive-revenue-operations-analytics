import json
import tempfile
import unittest
import os
from pathlib import Path
from unittest.mock import patch
from fastapi import FastAPI
from fastapi.testclient import TestClient
from docx import Document
from pptx import Presentation
from pypdf import PdfReader
from backend.tests.test_publications import sample
from backend.app.exporting import write_pack
from backend.app.reporting import generate,report_metadata
from backend.app.reports_api import create_router


class ReportingTests(unittest.TestCase):
    def test_normal_and_empty_artifacts_preserve_actual_kpis(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);exports=root/'exports';reports=root/'reports'
            for empty in [False,True]:
                source=write_pack(sample(empty),exports)
                result=generate(source.name,exports,reports);folder=reports/result['report_id']
                self.assertEqual(result['kpis']['new_leads'],0 if empty else 758)
                self.assertEqual(result['generation_id'],source.name)
                self.assertEqual(len(Presentation(folder/'LeadRevive_Management_SYNTHETIC.pptx').slides),10)
                text='\n'.join(p.text for p in Document(folder/'LeadRevive_Report_SYNTHETIC.docx').paragraphs)
                self.assertIn('External live delivery disabled',text)
                pages=[p.extract_text() for p in PdfReader(folder/'LeadRevive_Report_SYNTHETIC.pdf').pages]
                self.assertEqual(len(pages),2)
                self.assertNotIn('Automation and evidence',pages[0])
                self.assertIn('Automation and evidence',pages[1])
                self.assertIn('Evidence boundaries',pages[1])
                pdf=''.join(pages)
                self.assertIn(result['snapshot_id'],pdf)
                self.assertIn('N/A' if empty else '758',pdf)
                self.assertEqual(report_metadata(result['report_id'],reports)['data_revision'],result['data_revision'])

    def test_api_download_validation_and_atomic_failure(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);source=write_pack(sample(True),root/'exports')
            app=FastAPI();app.include_router(create_router(None,lambda:'2026-10-25T00:00:00Z',root/'exports',root/'reports'))
            with TestClient(app) as client:
                self.assertEqual(client.post('/api/reports/generate',json={'generation_id':'bad'}).status_code,422)
                with patch('backend.app.reporting.pdf_report',side_effect=RuntimeError('render failure')):
                    with self.assertRaises(RuntimeError):client.post('/api/reports/generate',json={'generation_id':source.name})
                self.assertEqual(client.get('/api/reports').json()['reports'],[])
                result=client.post('/api/reports/generate',json={'generation_id':source.name}).json()
                url=f"/api/reports/{result['report_id']}/download/LeadRevive_Report_SYNTHETIC.pdf"
                self.assertEqual(client.get(url).status_code,200)
                self.assertEqual(client.get(f"/api/reports/{result['report_id']}/download/.env").status_code,404)
                (root/'reports'/result['report_id']/'LeadRevive_Report_SYNTHETIC.pdf').write_text('tampered')
                self.assertEqual(client.get(url).status_code,404)

    @unittest.skipUnless(os.environ.get('LEADREVIVE_TEST_POSTGRES_URL'),'Fresh cohorts require PostgreSQL')
    def test_fresh_empty_postgres_cohort_does_not_replace_current(self):
        from backend.tests import test_foundation as ft
        ft.FoundationTests.setUp(self)
        try:
            ft.FoundationTests.ingest(self)
            root=self.root;source=write_pack(sample(True),root/'exports')
            app=FastAPI();app.include_router(create_router(self.engine,lambda:ft.NOW,root/'exports',root/'reports'))
            with TestClient(app) as client:
                response=client.post('/api/reports/generate',json=dict(generation_id=source.name,mode='fresh_cohort',start='2099-01-01T00:00:00Z',end='2100-01-01T00:00:00Z'))
                self.assertEqual(response.status_code,200,response.text)
                self.assertEqual(response.json()['kpis']['new_leads'],0)
                self.assertIsNone(response.json()['kpis']['median_first_response_hours'])
                from backend.app.artifacts import current_dir
                self.assertEqual(current_dir(root/'exports'),source)
        finally:ft.FoundationTests.tearDown(self)

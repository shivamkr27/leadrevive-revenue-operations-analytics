import copy
import json
import tempfile
import unittest
import threading
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
from backend.app.artifacts import current_dir,verified_snapshot,number_equal,revision
from backend.app.exporting import write_pack,fact_kpis
from backend.app.narrative import case_study

ROOT=Path(__file__).resolve().parents[2]


def sample(empty=False):
    s=json.loads((ROOT/'analytics/exports/snapshot.json').read_text(encoding='utf-8'))
    s.pop('power_bi_verification',None)
    if empty:
        s['facts']=[];s['sources']=[];s['owners']=[];s['stages']=[]
    s['kpis']=fact_kpis(s['facts']);s['sql_kpis']=dict(s['kpis'])
    return s


class PublicationTests(unittest.TestCase):
    def test_reader_sees_complete_old_generation_while_new_writer_is_staged(self):
        from backend.app.exporting import write_files
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);first=write_pack(sample(True),root)
            staged=threading.Event();release=threading.Event()
            def pause_after_write(s,path):
                write_files(s,path);staged.set()
                if not release.wait(15):raise TimeoutError('Test writer was not released')
            newer=sample(True);newer['experiment']['seed']=99
            with ThreadPoolExecutor(max_workers=1) as executor,patch('backend.app.exporting.write_files',side_effect=pause_after_write):
                task=executor.submit(write_pack,newer,root)
                try:
                    self.assertTrue(staged.wait(10))
                    visible,_=verified_snapshot(root)
                    self.assertEqual(visible,first)
                finally:release.set()
                second=task.result(timeout=10)
            self.assertEqual(verified_snapshot(root)[0],second)

    def test_narrative_uses_actual_empty_cohort_and_dynamic_findings(self):
        empty=sample(True);text=case_study(empty)
        self.assertIn('0 leads (0 seed records; 0 controlled fixtures)',text)
        self.assertNotIn('Referral has',text)
        self.assertIn('N/A',text)
        populated=sample();populated['sources'][0]['segment']='Changed source'
        populated['sources'][0]['overdue']=999
        self.assertIn('Changed source has the largest overdue count',case_study(populated))

    def test_empty_export_preserves_headers_and_nulls(self):
        with tempfile.TemporaryDirectory() as temp:
            path=write_pack(sample(True),Path(temp))
            _,s=verified_snapshot(Path(temp))
            self.assertIsNone(s['kpis']['median_first_response_hours'])
            self.assertTrue(all(r['status']=='PASS' for r in s['reconciliation']))
            self.assertFalse(number_equal(None,0))
            wb=load_workbook(path/'LeadRevive_Analysis_SYNTHETIC.xlsx',read_only=True)
            self.assertEqual(wb['LeadFacts'].max_row,1)
            self.assertEqual(wb['LeadFacts']['A1'].value,'lead_id');wb.close()

    def test_failed_generation_leaves_old_current_untouched(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);old=write_pack(sample(True),root)
            with patch('backend.app.exporting.write_files',side_effect=RuntimeError('forced interruption')):
                with self.assertRaises(RuntimeError):write_pack(sample(),root)
            self.assertEqual(current_dir(root),old)
            self.assertEqual(len(list((root/'generations').glob('.staging-*'))),0)
            verified_snapshot(root)

    def test_pinned_generation_and_stale_evidence(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);s=sample(True);first=write_pack(s,root)
            s['experiment'].pop('engine_verification',None)
            second=write_pack(s,root)
            self.assertNotEqual(first,second)
            self.assertTrue((first/'experiment_engine_verification.json').exists())
            self.assertFalse((second/'experiment_engine_verification.json').exists())
            self.assertEqual(current_dir(root,first.name),first)
            with self.assertRaises(ValueError):current_dir(root,'../escape')

    def test_mismatch_rejected_and_provenance_tracks_experiment(self):
        with tempfile.TemporaryDirectory() as temp:
            s=sample(True);s['sql_kpis']['contacted']=2
            with self.assertRaisesRegex(ValueError,'SQL KPI mismatch'):write_pack(s,Path(temp))
            self.assertFalse((Path(temp)/'current.json').exists())
        a=sample(True);b=copy.deepcopy(a);b['experiment']['seed']=999
        self.assertNotEqual(revision(a)[0],revision(b)[0])

    def test_actual_proof_required_and_tampering_detected(self):
        with tempfile.TemporaryDirectory() as temp:
            s=sample(True);data,provenance=revision(s)
            s['power_bi_verification']=dict(data_revision=data,model_sha256=provenance['model_sha256'],values=dict(s['kpis']))
            path=write_pack(s,Path(temp));_,published=verified_snapshot(Path(temp))
            self.assertTrue(all(r['power_bi_status'].startswith('PASS') for r in published['reconciliation']))
            (path/'KPI_SYNTHETIC.csv').write_text('tampered')
            with self.assertRaisesRegex(ValueError,'integrity failed'):verified_snapshot(Path(temp))
            s['power_bi_verification']['values']['median_first_response_hours']=0
            with self.assertRaisesRegex(ValueError,'Power BI KPI mismatch'):write_pack(s,Path(temp))

"""Null-aware export validation and atomic generation publication."""
import copy
import csv
import hashlib
import json
import shutil
import statistics
import tempfile
from pathlib import Path
from uuid import uuid4
from backend.app.analytics import csv_text,KPI_FIELDS,FACT_FIELDS
from backend.app.artifacts import revision,number_equal,atomic_json

SEGMENTS=['segment','leads','open','contacted','qualified','won','scheduled','overdue','pipeline_exposure_inr','sla_met','sla_breached','awaiting_within_sla','recovery_rule_eligible']
EXPERIMENT=['arm','cohort','surfaced','followups_completed','overdue_remaining','action_rate','timing_basis']
RULES=['arm','followups_completed','incomplete_due_work','raw_overdue_kpi','reminders_reserved','idempotent_hour_replays','measured_rule_execution_ms','outcome_basis']


def fact_kpis(rows):
    n=len(rows);tot=lambda key:sum(r[key] for r in rows)
    ratio=lambda numerator,denominator:numerator/denominator if denominator else None
    times=[r['first_response_hours'] for r in rows if r['first_response_hours'] is not None]
    return dict(new_leads=n,contacted=tot('contacted'),uncontacted=n-tot('contacted'),won=tot('won'),qualified=tot('qualified'),
        active_scheduled=tot('scheduled'),overdue=tot('overdue'),first_response_coverage=ratio(tot('contacted'),n),
        qualified_rate=ratio(tot('qualified'),n),win_rate=ratio(tot('won'),n),overdue_rate=ratio(tot('overdue'),tot('scheduled')),
        median_first_response_hours=statistics.median(times) if times else None,pipeline_exposure_inr=tot('pipeline_exposure_inr'))


def validate_values(expected,actual,label):
    for key in KPI_FIELDS:
        if key not in actual or not number_equal(expected[key],actual[key]):raise ValueError(f'{label} KPI mismatch: {key}')


def style_sheet(sheet,name,rows,fields):
    from openpyxl.styles import Font,PatternFill
    from openpyxl.worksheet.table import Table,TableStyleInfo
    sheet.append(fields)
    for row in rows:
        sheet.append([("'"+v if isinstance(v,str) and v.startswith(('=','+','-','@','\t','\r')) else v) for v in (row.get(f) for f in fields)])
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    if rows:
        table=Table(displayName='LR'+name,ref=sheet.dimensions)
        table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);sheet.add_table(table)
    for cell in sheet[1]:cell.fill=PatternFill('solid',fgColor='155E49');cell.font=Font(color='FFFFFF',bold=True)
    for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=min(55,max(14,max(len(str(c.value or '')) for c in col[:30])+2))
    sheet.print_title_rows='1:1'


def write_files(s,target):
    from openpyxl import Workbook,load_workbook
    from openpyxl.chart import BarChart,Reference
    validate_values(s['kpis'],s['sql_kpis'],'SQL')
    facts=fact_kpis(s['facts']);validate_values(s['kpis'],facts,'Fact aggregation')
    proof=s.get('power_bi_verification')
    if proof:
        if proof.get('data_revision')!=s['data_revision'] or proof.get('model_sha256')!=s['provenance']['model_sha256']:
            raise ValueError('Power BI proof is stale for this data/model revision')
        validate_values(s['kpis'],proof['values'],'Power BI')
    metadata={k:str(s[k]) for k in ['snapshot_id','generation_id','data_revision','origin','as_of_utc','window_start','window_end','source','interpretation']}
    metadata.update(model_sha256=s['provenance']['model_sha256'],application_version=s['provenance']['application_version'],
        experiment_assumptions=s['experiment']['assumptions'],operational_scope='All-time local test ledger including deliberate faults; not production reliability',
        composition=json.dumps({name:sum(r['cohort_kind']==name for r in s['facts']) for name in ['Seed dataset','Controlled fixture']}))
    tables={
        'LeadFacts':(s['facts'],list(FACT_FIELDS)),
        'KPI':([dict(metric=k,value=s['kpis'][k]) for k in KPI_FIELDS],['metric','value']),
        'Sources':(s['sources'],SEGMENTS),'Owners':(s['owners'],SEGMENTS),'Stages':(s['stages'],SEGMENTS),
        'Quality':(s['quality'],['reason','records','origin']),'Workflow':(s['workflow'],['workflow','status','jobs','attempts']),
        'Operations':([dict(metric=k,value=v) for k,v in s['operations'].items()],['metric','value']),
        'Experiment':(s['experiment']['summary'],EXPERIMENT),'ExperimentEvents':(s['experiment']['rows'],
            ['lead_id','due_hour','owner','manual_surface_hour','automation_surface_hour','manual_action_hour','automation_action_hour']),
        'ExperimentRules':(s['experiment'].get('engine_verification',[]),RULES),
        'ReadMe':([dict(field=k,value=v) for k,v in metadata.items()],['field','value'])}
    workbook=Workbook();workbook.remove(workbook.active)
    for name,(rows,fields) in tables.items():
        (target/(name+'_SYNTHETIC.csv')).write_text(csv_text(rows,fields),encoding='utf-8-sig',newline='')
        style_sheet(workbook.create_sheet(name),name,rows,fields)
    if s['stages']:
        chart=BarChart();chart.title='Current stage distribution | SYNTHETIC'
        chart.add_data(Reference(workbook['Stages'],min_col=2,min_row=1,max_row=1+len(s['stages'])),titles_from_data=True)
        chart.set_categories(Reference(workbook['Stages'],min_col=1,min_row=2,max_row=1+len(s['stages'])))
        workbook['Stages'].add_chart(chart,'R2')
    workbook.active=workbook.sheetnames.index('ReadMe')
    path=target/'LeadRevive_Analysis_SYNTHETIC.xlsx';workbook.save(path)
    with (target/'KPI_SYNTHETIC.csv').open(encoding='utf-8-sig',newline='') as stream:
        csv_values={r['metric']:float(r['value']) if r['value'] else None for r in csv.DictReader(stream)}
    check=load_workbook(path,read_only=True,data_only=True)
    excel=dict(check['KPI'].iter_rows(min_row=2,values_only=True));count=check['LeadFacts'].max_row-1;check.close()
    if count!=len(s['facts']):raise ValueError('Excel lead row count mismatch')
    validate_values(s['kpis'],csv_values,'CSV readback');validate_values(s['kpis'],excel,'Excel readback')
    reconciliation=[]
    for key in KPI_FIELDS:
        expected=s['kpis'][key];observed=[s['sql_kpis'][key],facts[key],csv_values[key],excel[key]]
        delta=max(abs(expected-v) for v in observed) if expected is not None else None
        passed=all(number_equal(expected,v) for v in observed)
        reconciliation.append(dict(metric=key,application=expected,sql=observed[0],facts=observed[1],csv=observed[2],excel=observed[3],
            difference=delta,status='PASS' if passed else 'FAIL',power_bi=proof['values'][key] if proof else None,
            power_bi_status='PASS: Desktop DAX verified' if proof and number_equal(expected,proof['values'][key]) else 'PENDING Desktop verification'))
    s['reconciliation']=reconciliation
    fields=list(reconciliation[0]);style_sheet(workbook.create_sheet('Reconciliation'),'Reconciliation',reconciliation,fields)
    (target/'Reconciliation_SYNTHETIC.csv').write_text(csv_text(reconciliation,fields),encoding='utf-8-sig',newline='')
    workbook.save(path)
    (target/'snapshot.json').write_text(json.dumps(s,indent=2),encoding='utf-8')
    if s['experiment'].get('engine_verification'):
        (target/'experiment_engine_verification.json').write_text(json.dumps(s['experiment'],indent=2),encoding='utf-8')
    manifest=dict(snapshot_id=s['snapshot_id'],generation_id=s['generation_id'],data_revision=s['data_revision'],provenance=s['provenance'],origin='SYNTHETIC',
        files={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in target.iterdir() if p.is_file()})
    (target/'manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')


def write_pack(snapshot,target,activate=True):
    target=Path(target);target.mkdir(parents=True,exist_ok=True)
    s=copy.deepcopy(snapshot);s.pop('reconciliation',None)
    s['data_revision'],s['provenance']=revision(s)
    if s.get('power_bi_verification') and 'data_revision' not in s['power_bi_verification']:
        s.pop('power_bi_verification') # Preserve legacy proof in the untouched Phase 4 files.
    s['generation_id']=uuid4().hex
    (target/'generations').mkdir(exist_ok=True)
    staging=Path(tempfile.mkdtemp(prefix='.staging-',dir=target/'generations'))
    try:
        write_files(s,staging)
        destination=target/'generations'/s['generation_id'];staging.rename(destination)
        if activate:atomic_json(target/'current.json',{'generation_id':s['generation_id'],'data_revision':s['data_revision']})
    finally:
        if staging.exists():shutil.rmtree(staging)
    return destination

"""Portable editable Word/PowerPoint and PDF, all from one verified generation."""
import hashlib
import json
import re
import shutil
import tempfile
from importlib.metadata import version
from datetime import datetime,timezone
from pathlib import Path
from uuid import uuid4
from xml.sax.saxutils import escape
from backend.app.artifacts import ROOT,EXPORTS,verified_snapshot
from backend.app.narrative import case_study,business_findings,cost_scenarios,LABELS,display

REPORTS=ROOT/'reports/generated'
IMAGES=ROOT/'reports/templates/evidence_images.json'


def sha(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def document_lines(s):
    """Compact cover provenance; complete identities remain in metadata.json."""
    for line in case_study(s).splitlines():
        if line.startswith('SYNTHETIC portfolio case | Snapshot '):
            yield f"SYNTHETIC portfolio case | Snapshot {s['snapshot_id']} | Revision {s['data_revision'][:12]}"
        elif line.startswith('Evaluation:'):
            scope='All stored creation dates through evaluation time' if s['window_start'].startswith('1900-') and s['window_end'].startswith('2100-') else f"Created from {s['window_start']} to {s['window_end']} (end excluded)"
            yield f"Evaluated {s['as_of_utc']}. {scope}. Current-state reporting."
        else:yield line


def evidence_images(snapshot):
    if not IMAGES.exists():return []
    manifest=json.loads(IMAGES.read_text(encoding='utf-8'))
    if manifest.get('data_revision')!=snapshot['data_revision']:return []
    result=[]
    for row in manifest['images']:
        path=ROOT/'screenshots'/row['file']
        if Path(row['file']).name!=row['file'] or sha(path)!=row['sha256']:raise ValueError('Screenshot evidence integrity failed')
        result.append(dict(row,path=str(path)))
    return result


def word_report(s,path,images):
    from docx import Document
    from docx.shared import Inches,Pt,RGBColor
    document=Document();normal=document.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10)
    normal.font.color.rgb=RGBColor.from_string('243D35')
    section=document.sections[0];section.top_margin=Inches(.65);section.bottom_margin=Inches(.65)
    header=section.header.paragraphs[0];header.text='LEADREVIVE  /  REVENUE OPERATIONS';header.style='Caption'
    for line in document_lines(s):
        if line.startswith('# '):document.add_heading(line[2:],0)
        elif line.startswith('## '):document.add_heading(line[3:],1)
        elif line.startswith('|---'):continue
        elif line.startswith('| '):
            cells=[c.strip() for c in line.strip('|').split('|')]
            if not locals().get('in_table',False):
                table=document.add_table(rows=0,cols=len(cells));table.style='Light Shading Accent 1'
            for cell,text in zip(table.add_row().cells,cells):cell.text=text
            in_table=True;continue
        elif line.startswith('- '):document.add_paragraph(line[2:],'List Bullet')
        elif line:document.add_paragraph(line)
        in_table=False
    if images:
        document.add_page_break();document.add_heading('Working interface evidence',1)
        for row in images[:3]:
            document.add_paragraph(row['caption']);document.add_picture(row['path'],width=Inches(6.1))
    footer=section.footer.paragraphs[0];footer.text=f"SYNTHETIC | {s['snapshot_id']} | {s['data_revision'][:12]} | External live delivery disabled"
    document.core_properties.title='LeadRevive — Business operations review';document.core_properties.subject=s['data_revision']
    document.save(path)


def pdf_report(s,path):
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,KeepTogether,PageBreak
    from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    styles=getSampleStyleSheet();styles['BodyText'].fontSize=9;styles['BodyText'].leading=13
    styles['BodyText'].spaceBefore=2
    styles['Heading1'].textColor=colors.HexColor('#155E49')
    styles['Heading1'].keepWithNext=True
    styles['Heading1'].spaceBefore=8
    story=[];table_rows=[]
    def flush():
        if not table_rows:return
        table=Table([[Paragraph(escape(c),styles['BodyText']) for c in row] for row in table_rows],repeatRows=1,hAlign='LEFT',colWidths=[290,200] if len(table_rows[0])==2 else [160,165,165])
        table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#E6F1EB')),('VALIGN',(0,0),(-1,-1),'TOP'),('BOTTOMPADDING',(0,0),(-1,-1),6),('LINEBELOW',(0,0),(-1,0),.5,colors.HexColor('#155E49'))]))
        story.extend([table,Spacer(1,8)]);table_rows.clear()
    for line in document_lines(s):
        if line.startswith('|---'):continue
        if line.startswith('| '):table_rows.append([c.strip() for c in line.strip('|').split('|')]);continue
        flush()
        if line.startswith('# '):story.append(Paragraph(escape(line[2:]),styles['Title']))
        elif line.startswith('## '):
            if line=='## Automation and evidence':story.append(PageBreak())
            story.append(Paragraph(escape(line[3:]),styles['Heading1']))
        elif line:story.extend([Paragraph(escape(line.lstrip('- ')),styles['BodyText']),Spacer(1,4)])
    flush()
    def footer(canvas,doc):
        canvas.saveState();canvas.setFont('Helvetica',7);canvas.setFillColor(colors.HexColor('#52665F'))
        canvas.drawString(44,24,f"SYNTHETIC | snapshot {s['snapshot_id']} | {s['data_revision'][:12]} | page {doc.page}");canvas.restoreState()
    SimpleDocTemplate(str(path),pagesize=A4,leftMargin=44,rightMargin=44,topMargin=40,bottomMargin=42,
        title='LeadRevive — Business operations review',author='LeadRevive portfolio project').build(story,onFirstPage=footer,onLaterPages=footer)


def slide_deck(s,path,images):
    from pptx import Presentation
    from pptx.util import Inches,Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from PIL import Image
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    green='155E49';dark='193C31';muted='52665F'
    def text(slide,value,x,y,w,h,size=20,color=dark,bold=False):
        box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));frame=box.text_frame;frame.word_wrap=True
        for index,line in enumerate(value.split('\n')):
            p=frame.paragraphs[0] if index==0 else frame.add_paragraph();p.text=line;p.font.name='Aptos';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(12)
        return box
    def page(title,subtitle='SYNTHETIC portfolio case'):
        slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string('F7FAF8')
        band=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,prs.slide_width,Inches(.12));band.fill.solid();band.fill.fore_color.rgb=RGBColor.from_string(green);band.line.fill.background()
        text(slide,'LEADREVIVE  /  MANAGEMENT REVIEW',.55,.3,12,.3,11,green,True)
        text(slide,title,.55,.85,12.1,.85,30,bold=True);text(slide,subtitle,.55,1.68,12.1,.55,13,muted)
        text(slide,f"SYNTHETIC  |  {s['snapshot_id']}  |  {s['data_revision'][:10]}  |  External live delivery disabled",.55,7.03,11.7,.3,9,muted)
        text(slide,f'{len(prs.slides):02}',12.1,7.03,.6,.3,10,muted)
        return slide
    def picture(slide,row,x=.6,y=2.3,w=8.2,h=4.45,crop=(.03,.36,.23,.10),center=False):
        left,right,top,bottom=crop
        with Image.open(row['path']) as im:ratio=(im.width*(1-left-right))/(im.height*(1-top-bottom))
        width=min(w,h*ratio);height=width/ratio
        if center:x+=(w-width)/2
        pic=slide.shapes.add_picture(row['path'],Inches(x),Inches(y),width=Inches(width),height=Inches(height))
        # Native PowerPoint viewport cropping only: original screenshot bytes stay intact.
        pic.crop_left=left;pic.crop_right=right;pic.crop_top=top;pic.crop_bottom=bottom
    def findings(slide,items):
        for index,value in enumerate(items):
            x=.65+index*4.18
            bar=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(5.94),Inches(3.85),Inches(.035))
            bar.fill.solid();bar.fill.fore_color.rgb=RGBColor.from_string(green);bar.line.fill.background()
            text(slide,value,x,6.07,3.85,.78,18,green,True)
    def image_by(token):return next((row for row in images if token in row['file']),None)
    k=s['kpis'];slide=page('Make missed follow-up work visible','Revenue leakage & lead recovery intelligence | Business operations portfolio')
    text(slide,'Question\nWhich enquiries need attention, who owns the next action, and can a repeatable process surface missed work?',.65,2.5,7,3,26)
    text(slide,f"{k['new_leads']:,}\nsynthetic records\n\n13 documented KPIs",9,2.5,3.5,3.5,27,green,True)
    slide=page('A reliable record before a reminder','Canonical PostgreSQL; normalized intake; conservative conflict review')
    text(slide,'Capture → validate → assign → prioritize → remind → review → measure',.65,2.4,12,1.1,25,green,True)
    text(slide,'Sales associate: own and reschedule work\nManager: assign overdue work and balance capacity\nAnalyst: reconcile cohorts, denominators and evidence',.65,3.8,11.5,2.4,22)
    slide=page('Where attention is needed',f"Frozen current-state cohort at {s['as_of_utc']}; not historical reconstruction")
    for index,(label,value) in enumerate([('Overdue / scheduled',f"{k['overdue']} / {k['active_scheduled']}"),('Estimated exposure',f"INR {k['pipeline_exposure_inr']:,.0f}"),('First-response median',display('median_first_response_hours',k['median_first_response_hours'])+' hours')]):
        text(slide,label,.65+index*4.2,2.5,3.9,.5,17,muted);text(slide,value,.65+index*4.2,3.15,3.9,1.2,29,green,True)
    text(slide,'Exposure is an estimated open-opportunity sum. It is not lost or recovered revenue. Uncontacted leads are excluded from the response-time median.',.65,5.25,11.8,1.2,20)
    slide=page('Diagnose workload before changing source spend')
    row=image_by('Analysis') or image_by('Owners')
    if row:
        picture(slide,row,x=.65,y=2.2,w=12,h=3.55,crop=(.035,.36,.40,.205),center=True)
        top_source=max(s['sources'],key=lambda r:r['overdue']) if s['sources'] else {'segment':'No source','overdue':0}
        unowned=sum(r['overdue'] for r in s['facts'] if r['owner_id']=='Unassigned')
        findings(slide,[f"{top_source['segment']}: {top_source['overdue']} overdue\nReview workload first",f"{unowned} overdue unassigned\nAssign the next owner",'Exposure is estimated\nNot recovered revenue'])
    else:text(slide,'\n\n'.join(business_findings(s)),.65,2.4,12,4.3,20)
    slide=page('Four working workflows, clear approval boundaries','Intake · recovery · outbox/approval · status sync')
    row=image_by('Workflows')
    if row:picture(slide,row,x=.65,y=2.2,w=12,h=3.55,crop=(.035,.36,.40,.215),center=True)
    else:text(slide,'All-time recorded local test ledger\n\n'+ '\n'.join(f"{key.replace('_',' ')}: {s['operations'].get(key,0)}" for key in ['distinct_reminded_leads','sink_only_approvals','smtp_accepted_by_local_sink']),.65,2.5,7.8,3.8,22)
    ops=s['operations']
    findings(slide,[f"{ops.get('distinct_reminded_leads',0)} leads surfaced\nOwner reminders reserved",f"{ops.get('sink_only_approvals',0)} sink-only approvals\nExact-content review",f"{ops.get('smtp_accepted_by_local_sink',0)} local sink acceptances\nExternal sending blocked"])
    slide=page('The comparison is modeled, not a sales result',s['experiment']['assumptions'])
    row=image_by('Simulation')
    if row:picture(slide,row,x=.65,y=2.2,w=12,h=3.55,crop=(.035,.36,.40,.225),center=True)
    else:text(slide,'\n'.join(f"{r['arm']}: {r['followups_completed']}/{r['cohort']} modeled contacts" for r in s['experiment']['summary']),.65,2.7,7.5,3,28)
    modeled={r['arm']:r for r in s['experiment']['summary']}
    findings(slide,[f"{modeled['manual']['followups_completed']} vs {modeled['automation']['followups_completed']} contacts\nModeled at hour 24",f"Same {modeled['manual']['cohort']} due events\nOwner timing assumed",'Raw overdue unchanged\nRescheduling still required'])
    slide=page('Trust depends on matching definitions and versions','SQL vs Python; CSV/Excel readback; Desktop DAX; pinned immutable generations')
    row=image_by('Reconciliation')
    if row:
        picture(slide,row,x=.65,y=2.25,w=12,h=3.45,crop=(.04,.57,.31,.505),center=True)
        findings(slide,['13 KPIs reconciled\nSelected rows enlarged','SQL / exports / DAX\nSame snapshot and revision','Null is different from zero\nFull detail in the report'])
    else:text(slide,'No matching screenshot evidence for this cohort.\nThe report still includes all actual KPI values and source hashes.',.65,2.7,12,3,24)
    slide=page('Cost / benefit is a sensitivity exercise','Assumed capacity value, not observed cash savings or ROI')
    scenarios=cost_scenarios();rows=[['Assumed min saved/action','Capacity hours/month','Net capacity-equivalent INR/month']]+[[str(r['assumed_minutes_saved_per_action']),str(r['modeled_capacity_hours']),f"{r['net_capacity_equivalent_inr']:,.0f}"] for r in scenarios['scenarios']]
    table=slide.shapes.add_table(len(rows),3,Inches(.65),Inches(2.5),Inches(12),Inches(2)).table
    for i,row in enumerate(rows):
        for j,value in enumerate(row):
            cell=table.cell(i,j);cell.text=value;cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string('E6F1EB' if i==0 else 'FFFFFF')
            for p in cell.text_frame.paragraphs:p.font.size=Pt(17);p.font.color.rgb=RGBColor.from_string(dark)
    text(slide,'Assumptions: 200 actions/month; INR 300/hour; 24 setup hours; 4 maintenance hours/month; 12-month amortization. Hardware/electricity/licensing not measured.',.65,5.15,12,1.3,18)
    slide=page('What a permissioned pilot should measure')
    text(slide,'1  Assign ownership and require a next action after contact.\n2  Measure owner response and actual administrative effort.\n3  Separate induced test faults from ordinary operations.\n4  Track completed follow-up before making revenue claims.',.65,2.5,11.9,3.8,23)
    slide=page('Review gate and limitations','Local single-user prototype; no public deployment or external email')
    text(slide,'Synthetic data and retained test fixtures\nModel assumptions are not observed business outcomes\nCurrent-state snapshots are not time-travel reporting\nTested duplicate protection is not a universal exactly-once guarantee\n\nPhase 6, publishing and final sign-off require separate approval.',.65,2.5,11.9,4.1,22)
    for slide in prs.slides:
        slide.notes_slide.notes_text_frame.text=f"Evidence: snapshot {s['snapshot_id']}; data revision {s['data_revision']}; export generation {s['generation_id']}. All data synthetic. Figures use verified KPI fields; image hashes recorded in metadata.json. Original Power BI window screenshots are cropped to the report canvas in slide viewports only; original pixels/files are retained. No recovered revenue or external delivery claim."
    prs.core_properties.title='LeadRevive — Revenue operations management review';prs.core_properties.subject=s['data_revision'];prs.save(path)


def generate(generation=None,exports=EXPORTS,report_root=REPORTS):
    directory,s=verified_snapshot(exports,generation)
    if not s.get('generation_id'):raise ValueError('Publish a validated analytics generation first')
    report_root=Path(report_root);report_root.mkdir(parents=True,exist_ok=True)
    staging=Path(tempfile.mkdtemp(prefix='.staging-',dir=report_root));report_id=uuid4().hex
    try:
        images=evidence_images(s)
        (staging/'Case_Study_SYNTHETIC.md').write_text(case_study(s),encoding='utf-8')
        word_report(s,staging/'LeadRevive_Report_SYNTHETIC.docx',images)
        pdf_report(s,staging/'LeadRevive_Report_SYNTHETIC.pdf')
        slide_deck(s,staging/'LeadRevive_Management_SYNTHETIC.pptx',images)
        (staging/'Cost_Benefit_Assumptions.json').write_text(json.dumps(cost_scenarios(),indent=2),encoding='utf-8')
        metadata=dict(report_id=report_id,snapshot_id=s['snapshot_id'],generation_id=s['generation_id'],data_revision=s['data_revision'],
            origin='SYNTHETIC',as_of_utc=s['as_of_utc'],window_start=s['window_start'],window_end=s['window_end'],
            created_at_utc=datetime.now(timezone.utc).isoformat(),generator_sha256=sha(__file__),
            narrative_sha256=sha(Path(__file__).with_name('narrative.py')),
            libraries={name:version(name) for name in ['python-pptx','python-docx','reportlab']},
            source_manifest_sha256=sha(directory/'manifest.json'),kpis=s['kpis'],
            image_evidence=[{k:v for k,v in row.items() if k!='path'} for row in images],
            files={p.name:sha(p) for p in staging.iterdir() if p.is_file()},
            interpretation='Synthetic snapshot. Modeled outcomes are assumptions; SMTP accepted by local sink only. External live delivery disabled.')
        (staging/'metadata.json').write_text(json.dumps(metadata,indent=2),encoding='utf-8')
        destination=report_root/report_id;staging.rename(destination)
        return metadata
    finally:
        if staging.exists():shutil.rmtree(staging)


def report_metadata(report_id,root=REPORTS):
    if not re.fullmatch(r'[a-f0-9]{32}',report_id):raise ValueError('Invalid report identifier')
    path=Path(root)/report_id/'metadata.json'
    result=json.loads(path.read_text(encoding='utf-8'))
    if result['report_id']!=report_id:raise ValueError('Report identity mismatch')
    return result

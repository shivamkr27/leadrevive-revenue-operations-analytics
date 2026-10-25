"""Read-only verification of exports, query pack, API and preserved seed rows."""
import csv
import hashlib
import json
import sys
from pathlib import Path
import httpx
from openpyxl import load_workbook
from sqlalchemy import select,inspect,text

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from backend.app import models as m
from backend.app.foundation import engine_for
from scripts.local_config import postgres_url
from scripts.migrate_to_postgres import digest
from backend.app.artifacts import current_dir,number_equal


def run():
    out=current_dir()
    d=json.loads((out/'snapshot.json').read_text())
    manifest=json.loads((out/'manifest.json').read_text())
    for name,expected in manifest['files'].items():
        assert hashlib.sha256((out/name).read_bytes()).hexdigest()==expected,name
    with (out/'KPI_SYNTHETIC.csv').open(encoding='utf-8-sig',newline='') as stream:
        csv_values={r['metric']:float(r['value']) if r['value'] else None for r in csv.DictReader(stream)}
    workbook=load_workbook(out/'LeadRevive_Analysis_SYNTHETIC.xlsx',read_only=True,data_only=True)
    excel_values=dict(workbook['KPI'].iter_rows(min_row=2,values_only=True))
    for key,value in d['kpis'].items():
        for comparison in [csv_values[key],excel_values[key],d['sql_kpis'][key],d['power_bi_verification']['values'][key]]:
            assert number_equal(value,comparison),(key,value,comparison)
    assert workbook['LeadFacts'].max_row-1==d['kpis']['new_leads']
    workbook.close()
    response=httpx.get('http://127.0.0.1:8000/api/analytics/report');response.raise_for_status()
    app=response.json()
    assert app['snapshot_id']==d['snapshot_id'] and app['kpis']==d['kpis']
    db=engine_for(postgres_url());old=engine_for('sqlite:///'+(ROOT/'data/processed/leadrevive.sqlite').as_posix())
    preserved={};queries=0
    with old.connect() as before,db.connect() as current:
        for table in m.metadata.sorted_tables:
            if table.name not in inspect(old).get_table_names():continue
            rows=[dict(r) for r in before.execute(select(table)).mappings()]
            key=list(table.primary_key.columns)[0].name
            identifiers={r[key] for r in rows}
            after=[dict(r) for r in current.execute(select(table)).mappings() if r[key] in identifiers]
            assert digest(rows)==digest(after),table.name
            preserved[table.name]=len(rows)
        for name in ['management_analysis.sql','automation_evidence.sql']:
            sql='\n'.join(line.split('--')[0] for line in (ROOT/'analytics/sql'/name).read_text().splitlines())
            for query in sql.split(';'):
                if query.strip():
                    current.execute(text(query),dict(as_of=d['as_of_utc'],cooldown_hours=24)).all()
                    queries+=1
    old.dispose();db.dispose()
    result=dict(result='passed',snapshot_id=d['snapshot_id'],reconciled_kpis=len(d['kpis']),
                manifest_files=len(manifest['files']),executed_management_queries=queries,
                original_rows_preserved=preserved,power_bi_method=d['power_bi_verification']['method'])
    (ROOT/'data/processed/phase4_pack_verification.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
    print(json.dumps(result,indent=2))


if __name__=='__main__':run()
